"""Сервис для генерации Excel отчетов.

Модуль содержит функции для генерации Excel файлов с аналитикой:
- Почасовой график запусков бота новыми пользователями
- Почасовой график оплат
- Почасовой график пользователей, пришедших по реферальной ссылке
- Таблица пользователей с количеством рефералов
"""
import os
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_
from app.models import User, Payment

logger = logging.getLogger(__name__)


async def generate_excel_report(
    session: AsyncSession,
    start_date: datetime,
    end_date: datetime,
    output_path: str
) -> str:
    """Генерирует Excel отчет с аналитикой за указанный период.
    
    Создает Excel файл с двумя листами:
    1. Почасовой график с метриками:
       - Количество запусков бота новыми пользователями
       - Количество оплат
       - Количество пользователей, пришедших по реферальной ссылке
    2. Таблица пользователей с количеством рефералов
    
    Args:
        session (AsyncSession): Асинхронная сессия БД
        start_date (datetime): Начало периода
        end_date (datetime): Конец периода (включительно)
        output_path (str): Путь для сохранения файла
        
    Returns:
        str: Путь к созданному файлу
    """
    try:
        # Создаем рабочую книгу
        wb = Workbook()
        
        # Удаляем дефолтный лист
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])
        
        # Создаем лист с почасовым графиком
        logger.info(f"Creating hourly chart sheet")
        hourly_sheet = wb.create_sheet("Почасовой график")
        await _create_hourly_chart(session, hourly_sheet, start_date, end_date)
        
        # Создаем лист с таблицей рефералов
        logger.info(f"Creating referrals table sheet")
        referrals_sheet = wb.create_sheet("Рефералы")
        await _create_referrals_table(session, referrals_sheet)
        
        # Сохраняем файл
        logger.info(f"Saving Excel file to {output_path}")
        wb.save(output_path)
        
        # Проверяем, что файл создан
        if not os.path.exists(output_path):
            raise FileNotFoundError(f"Excel file was not created at {output_path}")
        
        logger.info(f"Excel report generated successfully: {output_path}")
        return output_path
    except Exception as e:
        logger.error(f"Error in generate_excel_report: {e}", exc_info=True)
        raise


async def _create_hourly_chart(
    session: AsyncSession,
    sheet,
    start_date: datetime,
    end_date: datetime
) -> None:
    """Создает почасовой график метрик.
    
    Args:
        session (AsyncSession): Асинхронная сессия БД
        sheet: Лист Excel для записи данных
        start_date (datetime): Начало периода
        end_date (datetime): Конец периода
    """
    # Заголовки
    headers = [
        "Дата и время",
        "Запуски бота новыми пользователями",
        "Количество оплат",
        "Пользователи по реферальной ссылке"
    ]
    
    # Записываем заголовки
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Генерируем список всех часов в периоде
    # Сохраняем timezone если он есть
    if start_date.tzinfo:
        current = start_date.replace(minute=0, second=0, microsecond=0)
        end = end_date.replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
    else:
        # Если timezone нет, используем UTC
        from datetime import timezone as tz
        current = start_date.replace(minute=0, second=0, microsecond=0, tzinfo=tz.utc)
        end = (end_date.replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)).replace(tzinfo=tz.utc)
    
    row = 2
    while current < end:
        # Получаем данные за этот час
        hour_start = current
        hour_end = current + timedelta(hours=1)
        
        # Количество запусков бота новыми пользователями (созданные в этот час)
        new_users_query = select(func.count(User.user_id)).where(
            and_(
                User.created_at >= hour_start,
                User.created_at < hour_end
            )
        )
        new_users_result = await session.execute(new_users_query)
        new_users_count = new_users_result.scalar() or 0
        
        # Количество оплат за этот час
        payments_query = select(func.count(Payment.id)).where(
            and_(
                Payment.created_at >= hour_start,
                Payment.created_at < hour_end,
                Payment.status == "completed"
            )
        )
        payments_result = await session.execute(payments_query)
        payments_count = payments_result.scalar() or 0
        
        # Количество пользователей, пришедших по реферальной ссылке (с referrer_id)
        referrals_query = select(func.count(User.user_id)).where(
            and_(
                User.created_at >= hour_start,
                User.created_at < hour_end,
                User.referrer_id.isnot(None)
            )
        )
        referrals_result = await session.execute(referrals_query)
        referrals_count = referrals_result.scalar() or 0
        
        # Записываем данные
        sheet.cell(row=row, column=1, value=hour_start.strftime("%d.%m.%Y %H:00"))
        sheet.cell(row=row, column=2, value=new_users_count)
        sheet.cell(row=row, column=3, value=payments_count)
        sheet.cell(row=row, column=4, value=referrals_count)
        
        current += timedelta(hours=1)
        row += 1
    
    # Настраиваем ширину колонок
    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 35
    sheet.column_dimensions['C'].width = 20
    sheet.column_dimensions['D'].width = 35
    
    # Замораживаем первую строку
    sheet.freeze_panes = 'A2'


async def _create_referrals_table(session: AsyncSession, sheet) -> None:
    """Создает таблицу пользователей с количеством рефералов.
    
    Args:
        session (AsyncSession): Асинхронная сессия БД
        sheet: Лист Excel для записи данных
    """
    # Заголовки
    headers = ["Пользователь (user_id)", "Количество рефералов"]
    
    # Записываем заголовки
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Получаем всех пользователей с рефералами
    users_query = select(User).where(
        User.referrals.isnot(None)
    )
    users_result = await session.execute(users_query)
    users = users_result.scalars().all()
    
    # Сортируем по количеству рефералов (по убыванию)
    users_with_refs = [
        (user.user_id, len(user.referrals) if user.referrals else 0)
        for user in users
        if user.referrals and len(user.referrals) > 0
    ]
    users_with_refs.sort(key=lambda x: x[1], reverse=True)
    
    # Записываем данные
    row = 2
    for user_id, ref_count in users_with_refs:
        sheet.cell(row=row, column=1, value=user_id)
        sheet.cell(row=row, column=2, value=ref_count)
        row += 1
    
    # Настраиваем ширину колонок
    sheet.column_dimensions['A'].width = 25
    sheet.column_dimensions['B'].width = 25
    
    # Замораживаем первую строку
    sheet.freeze_panes = 'A2'

